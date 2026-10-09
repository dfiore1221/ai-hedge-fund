from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "AIFundOS_Infrastructure_and_Strategy_Roadmap.docx"

INK = "000000"
NAVY = "17365D"
MID_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
PALE_GRAY = "F5F5F5"
GRID = "D9D9D9"
MUTED = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=130, bottom=120, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    node.set(qn("w:val"), "1" if value else "0")


def set_repeat_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("AIFundOS Infrastructure and Strategy Roadmap")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AIFundOS Roadmap  |  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    set_repeat_header_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor(31, 31, 31)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 13.5, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.left_indent = Inches(0.26)
        style.paragraph_format.first_line_indent = Inches(-0.16)

    if "Checklist" not in styles:
        checklist = styles.add_style("Checklist", WD_STYLE_TYPE.PARAGRAPH)
        checklist.base_style = styles["Normal"]
        checklist.font.name = "Aptos"
        checklist.font.size = Pt(10.5)
        checklist.paragraph_format.left_indent = Inches(0.24)
        checklist.paragraph_format.first_line_indent = Inches(-0.24)
        checklist.paragraph_format.space_after = Pt(4)


def add_bullet(doc, text, style="List Bullet"):
    return doc.add_paragraph(text, style=style)


def add_checklist(doc, items):
    for item in items:
        doc.add_paragraph(f"[ ]  {item}", style="Checklist")


def set_col_width(cell, width):
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_summary_table(doc):
    rows = [
        ("1", "Always-On Host", "One authoritative worker that stays online", "Not started"),
        ("2", "Streaming Equity Data", "Fresh quotes, bid and ask, one-minute bars", "Not started"),
        ("3", "Event-Driven Execution", "Durable order, fill, position, and recovery flow", "Not started"),
        ("4", "Strategy Data", "Options and short-selling data with realistic costs", "Blocked on data"),
        ("5", "Staged Strategy Rollout", "Expand only after infrastructure and evidence gates", "Swing active"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.5, 1.55, 3.25, 1.25]
    headers = ["Phase", "Workstream", "Outcome", "Status"]
    for idx, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        set_col_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for row_idx, row_data in enumerate(rows, start=1):
        cells = table.add_row().cells
        for idx, (cell, text) in enumerate(zip(cells, row_data)):
            set_col_width(cell, widths[idx])
            set_cell_shading(cell, "FFFFFF" if row_idx % 2 else PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = "Aptos"
            run.font.size = Pt(9.3)
            if idx == 1:
                run.bold = True
    set_table_borders(table)
    doc.add_paragraph()


def add_phase(doc, number, title, purpose, checklist, done, extra_sections=None):
    heading = doc.add_heading(f"Phase {number}  {title}", level=1)
    set_keep_with_next(heading)
    p = doc.add_paragraph()
    lead = p.add_run("Purpose  ")
    lead.bold = True
    p.add_run(purpose)

    doc.add_heading("Required Work", level=2)
    add_checklist(doc, checklist)

    if extra_sections:
        for subheading, paragraphs in extra_sections:
            doc.add_heading(subheading, level=2)
            for para in paragraphs:
                if isinstance(para, tuple) and para[0] == "check":
                    add_checklist(doc, para[1])
                else:
                    doc.add_paragraph(para)

    doc.add_heading("Definition of Done", level=2)
    for item in done:
        add_bullet(doc, item)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("AIFundOS Infrastructure and Strategy Roadmap")
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("Always-on operations, market data, execution, and strategy expansion")
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.color.rgb = MUTED
    meta = doc.add_paragraph()
    meta_run = meta.add_run("Updated October 8, 2026")
    meta_run.font.name = "Aptos"
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = MUTED
    meta.paragraph_format.space_after = Pt(18)

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This roadmap moves AIFundOS from a laptop-dependent autonomous paper-trading system "
        "to an always-on, event-driven research and paper-execution platform. The immediate "
        "priority is dependable uptime and execution-quality equity data. Faster or more complex "
        "strategies should be introduced only after the supporting infrastructure, data, and "
        "evidence gates are working reliably."
    )
    doc.add_paragraph(
        "The main conclusion is straightforward: an always-on host solves availability, while a "
        "streaming feed and event-driven execution engine solve market awareness and order "
        "realism. All three are needed before AIFundOS can responsibly test day trading, options, "
        "short selling, or scalping at meaningful scale."
    )

    doc.add_heading("Roadmap at a Glance", level=1)
    add_summary_table(doc)

    doc.add_heading("Current State", level=1)
    doc.add_heading("Operational Today", level=2)
    for item in (
        "Autonomous paper-only swing and position workflow.",
        "Morning brief, risk review, portfolio governor, trade journal, paper ledger, memory, email alerts, and end-of-day review.",
        "Desktop watchdog every 10 minutes during the trading session and paper entry or exit checks every 5 minutes.",
        "Tiingo authenticated equity data with Yahoo fallback.",
        "Benzinga, Finnhub, FRED, SEC EDGAR, and Quiver integrations.",
        "Docker worker and dashboard deployment package.",
        "Data-quality gate currently passes for the swing workflow.",
    ):
        add_bullet(doc, item)

    doc.add_heading("Important Limitations", level=2)
    for item in (
        "The Mac laptop remains the active host and stops working when it is asleep, offline, or traveling.",
        "Equity data is polled rather than consumed as a continuous streaming feed.",
        "Provider checks can fall back to prior closes or cached observations.",
        "Options data does not provide dependable OPRA-quality quotes, historical Greeks, implied-volatility history, or options flow.",
        "No borrow, locate, short-sale restriction, margin, or borrow-fee model exists for short selling.",
        "No continuous event engine supervises same-session strategies.",
        "Costs and lifecycle handling remain incomplete for dividends, interest, borrow fees, assignment, and exercise.",
        "Tactical evidence remains below the 30-completed-trade minimum gate.",
    ):
        add_bullet(doc, item)

    add_phase(
        doc,
        1,
        "Always-On Private Host",
        "Run one authoritative AIFundOS worker continuously on a wired Linux mini PC, Mac mini, or private cloud host.",
        [
            "Choose the host and connect it through wired Ethernet.",
            "Add UPS battery protection for a home host.",
            "Deploy the existing Docker Compose worker and dashboard.",
            "Enable automatic restart after a crash, reboot, or power interruption.",
            "Use Tailscale or another private authenticated tunnel; do not expose Streamlit directly to the public internet.",
            "Store secrets securely and enable encrypted backups.",
            "Persist the portfolio ledger, memory database, reports, and data cache.",
            "Synchronize system time and enforce the America New York market clock.",
            "Alert on missed runs, stale files, failed email, provider failures, and disk pressure.",
            "Permit exactly one ledger-writing worker; all dashboard replicas remain read-only.",
        ],
        [
            "Twenty consecutive U.S. market sessions with at least 99.5 percent scheduled-service availability.",
            "No missed morning brief, market-session worker, end-of-day review, or Friday review caused by host sleep or Wi-Fi loss.",
            "Successful restart recovery without duplicate fills or ledger divergence.",
            "A tested encrypted backup and restore of portfolio and memory data.",
        ],
    )

    add_phase(
        doc,
        2,
        "Streaming Equity Data",
        "Give AIFundOS continuous, timestamped market awareness and retain an independent provider for validation.",
        [
            "Select a primary provider with WebSocket streaming quotes.",
            "Normalize symbol, exchange, event time, receive time, and sequence identifiers where available.",
            "Capture last trade, bid, ask, spread, size, cumulative volume, and session status.",
            "Build one-minute OHLCV bars and retain prior adjusted close.",
            "Capture split, dividend, symbol-change, and trading-halt context.",
            "Reject stale and out-of-order quotes.",
            "Reject execution when bid, ask, or timestamps are missing.",
            "Detect material disagreement between primary and validation providers.",
            "Persist normalized events so decisions can be replayed and audited.",
            "Separate regular-session, extended-hours, and closed-market observations.",
        ],
        [
            "At least 99.9 percent complete one-minute bars during regular market sessions for the active universe.",
            "Quote age and provider status are visible on every execution decision.",
            "Twenty sessions without fills based on stale, cached, or prior-close data.",
        ],
    )

    add_phase(
        doc,
        3,
        "Event-Driven Paper Execution",
        "Replace interval-only checks with a durable pipeline from market event through trade review.",
        [
            "Create the pipeline: market event to candidate update to strategy gate to risk gate to working order to simulated fill to position supervision to exit to review.",
            "Assign idempotent order and fill identifiers.",
            "Prevent duplicate orders across crashes and restarts.",
            "Add a gap-through-entry veto and mandatory re-underwriting after abnormal adverse gaps.",
            "Model spread, slippage, liquidity, and partial fills.",
            "Define behavior for trading halts and missing data.",
            "Force same-day liquidation for day strategies.",
            "Recover cleanly after restarts and reconcile the ledger.",
            "Track the full funnel from discovery through realized profit or loss.",
        ],
        [
            "Zero duplicate fills across restart tests.",
            "Every fill is traceable to the exact quote, strategy decision, risk decision, and working order.",
            "Replay tests produce the same ledger result from the same event stream.",
        ],
    )

    add_phase(
        doc,
        4,
        "Strategy-Specific Data",
        "Add the data and lifecycle controls required to test options and short selling realistically.",
        [
            "Options: connect OPRA-quality chains with bid, ask, Greeks, implied volatility, historical implied volatility, volume, open interest, and contract adjustments.",
            "Options: model conservative multi-leg synchronization, assignment, exercise, expiration, and spread slippage.",
            "Shorts: connect borrow availability, hard-to-borrow status, estimated borrow cost, short-sale restriction status, and margin requirements.",
            "Shorts: model dividends owed, recall risk, and strict gap-loss controls.",
        ],
        [
            "No options strategy is enabled until quote, Greeks, lifecycle, and cost checks pass.",
            "No short strategy is enabled until borrow, margin, restriction, and fee checks pass.",
            "Paper results include realistic transaction and carrying costs.",
        ],
    )

    add_phase(
        doc,
        5,
        "Staged Strategy Introduction",
        "Expand the Committee's trading freedom in controlled stages that generate useful evidence without producing unrealistic paper results.",
        [
            "Continue swing and position paper experiments under current risk constraints.",
            "Add equity-short signals in shadow mode with no fills.",
            "Enable small defined-risk paper shorts after the short-data gate passes.",
            "Enable long calls and long puts after the options-data and lifecycle gates pass.",
            "Enable day strategies only after 20 reliable always-on sessions and event-engine verification.",
            "Consider scalping last and keep it disabled unless sub-minute supervision and realistic spread and slippage evidence justify it.",
        ],
        [
            "Each strategy family has enough completed observations to estimate expectancy after costs.",
            "Risk increases are based on evidence, not simply on a strategy becoming technically available.",
            "Underperforming strategy families can be paused independently without disrupting the rest of AIFundOS.",
        ],
    )

    doc.add_heading("Evidence and Risk Gates", level=1)
    doc.add_paragraph(
        "Infrastructure readiness and investment edge are separate questions. Reliable uptime does not prove a strategy is profitable, and promising backtests do not make an unreliable execution system safe. Both gates must pass before AIFundOS increases autonomy or risk."
    )
    for item in (
        "Keep tactical risk constrained until at least 30 autonomous tactical trades close.",
        "Evaluate expectancy after spread, slippage, and modeled costs.",
        "Track win rate, average win, average loss, drawdown, profit factor, exposure, and benchmark-relative return by strategy family.",
        "Require out-of-sample or walk-forward evidence before increasing size.",
        "Block any strategy that lacks its required data even when the overall data-quality score is high.",
    ):
        add_bullet(doc, item)

    doc.add_heading("Recommended Build Order", level=1)
    for item in (
        "Choose and provision the always-on host.",
        "Deploy the existing Docker worker and dashboard privately.",
        "Prove uptime and single-writer ledger safety.",
        "Select and integrate the streaming equity provider.",
        "Build normalized market-event storage and freshness gates.",
        "Implement the event-driven paper-execution engine.",
        "Add short-selling data, followed by options data and lifecycle handling.",
    ):
        add_bullet(doc, item, style="List Number")

    doc.add_heading("Immediate Next Decision", level=1)
    doc.add_paragraph(
        "Choose the always-on host. A wired mini PC or Mac mini at home is the most direct path if the goal is private local control. A private cloud host is easier to keep online but requires ongoing hosting cost and careful security. Once the host is chosen, deploy the existing Docker package and begin the 20-session reliability test before adding faster strategies."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
